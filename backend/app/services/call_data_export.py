"""
Call Data Export Service - Export greeting and call data to Excel
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.db.models.call_session import CallSession
from app.db.models.conversation import Conversation, ConversationMessage
from app.core.exceptions import NotFoundException

logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed. Excel export will be unavailable.")


class CallDataExporter:
    """Export greeting and call data to Excel files"""
    
    def __init__(self, db: Session, output_dir: str = "exports") -> None:
        self.db = db
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
        
        if not OPENPYXL_AVAILABLE:
            logger.warning("Excel export not available: openpyxl not installed")

    def export_greeting_data(
        self,
        call_session_id: int,
        filename: str | None = None
    ) -> Path | None:
        """
        Export greeting and call data to Excel
        
        Includes:
        - Call info (ID, caller number, date/time)
        - Greeting status (initial, language selected, confirmation)
        - Language selection
        - Timestamp data
        """
        if not OPENPYXL_AVAILABLE:
            logger.error("Excel export not available: openpyxl not installed")
            return None
        
        try:
            call = self.db.query(CallSession).filter(
                CallSession.id == call_session_id
            ).first()
            
            if not call:
                raise NotFoundException(f"Call session {call_session_id} not found")
            
            conversation = self.db.query(Conversation).filter(
                Conversation.call_session_id == call_session_id
            ).first()
            
            # Create workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Greeting Data"
            
            # Style definitions
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )
            
            # Set column widths
            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 40
            
            # Title
            ws["A1"] = "GREETING & CALL DATA"
            ws["A1"].font = Font(bold=True, size=14)
            ws.merge_cells("A1:B1")
            
            # Section 1: Call Information
            row = 3
            self._add_section_header(ws, row, "CALL INFORMATION", header_fill, header_font, border)
            row += 1
            
            self._add_data_row(ws, row, "Call ID:", call.id, border)
            row += 1
            self._add_data_row(ws, row, "Caller Number:", call.caller_number, border)
            row += 1
            self._add_data_row(ws, row, "Caller Name:", call.caller_name, border)
            row += 1
            self._add_data_row(ws, row, "Device ID:", call.device_id, border)
            row += 1
            self._add_data_row(ws, row, "Call Status:", call.call_status, border)
            row += 1
            self._add_data_row(ws, row, "Call Date/Time:", call.created_at.isoformat() if call.created_at else "N/A", border)
            row += 1
            
            # Section 2: Greeting Status
            if conversation:
                row += 1
                self._add_section_header(ws, row, "GREETING STATUS", header_fill, header_font, border)
                row += 1
                
                self._add_data_row(ws, row, "Conversation ID:", conversation.id, border)
                row += 1
                self._add_data_row(ws, row, "Status:", conversation.status, border)
                row += 1
                self._add_data_row(ws, row, "Language Selected:", conversation.language or "Not selected", border)
                row += 1
                self._add_data_row(ws, row, "Conversation Started:", conversation.started_at.isoformat() if conversation.started_at else "N/A", border)
                row += 1
                self._add_data_row(ws, row, "Conversation Completed:", conversation.completed_at.isoformat() if conversation.completed_at else "Not completed", border)
                row += 1
                
                # Section 3: Greeting Messages
                greeting_messages = [
                    msg for msg in conversation.messages
                    if msg.message_type in ["greeting", "language_prompt", "language_confirmation"]
                ]
                
                if greeting_messages:
                    row += 1
                    self._add_section_header(ws, row, "GREETING MESSAGES", header_fill, header_font, border)
                    row += 1
                    
                    for idx, msg in enumerate(greeting_messages, 1):
                        self._add_data_row(ws, row, f"Message {idx}:", "", border)
                        row += 1
                        
                        ws[f"A{row}"] = "  Type:"
                        ws[f"B{row}"] = msg.message_type
                        ws[f"A{row}"].border = border
                        ws[f"B{row}"].border = border
                        row += 1
                        
                        ws[f"A{row}"] = "  Language:"
                        ws[f"B{row}"] = msg.language or "Not specified"
                        ws[f"A{row}"].border = border
                        ws[f"B{row}"].border = border
                        row += 1
                        
                        ws[f"A{row}"] = "  Content:"
                        ws[f"B{row}"] = msg.content[:100] + "..." if len(msg.content) > 100 else msg.content
                        ws[f"A{row}"].border = border
                        ws[f"B{row}"].border = border
                        ws[f"B{row}"].alignment = Alignment(wrap_text=True)
                        row += 1
                        
                        ws[f"A{row}"] = "  Timestamp:"
                        ws[f"B{row}"] = msg.created_at.isoformat() if msg.created_at else "N/A"
                        ws[f"A{row}"].border = border
                        ws[f"B{row}"].border = border
                        row += 1
                        row += 1
                
                # Section 4: Summary
                row += 1
                self._add_section_header(ws, row, "SUMMARY", header_fill, header_font, border)
                row += 1
                
                greeting_played = any(
                    msg.message_type == "greeting" 
                    for msg in conversation.messages
                )
                language_selected = conversation.language is not None
                
                self._add_data_row(ws, row, "Greeting Played:", "Yes" if greeting_played else "No", border)
                row += 1
                self._add_data_row(ws, row, "Language Selected:", "Yes" if language_selected else "No", border)
                row += 1
                self._add_data_row(ws, row, "Total Questions:", conversation.total_questions, border)
                row += 1
                self._add_data_row(ws, row, "Current Question:", conversation.current_question_index, border)
                row += 1
                self._add_data_row(ws, row, "Completion %:", f"{conversation.completion_percentage}%", border)
                row += 1
                
                # Export timestamp
                row += 1
                self._add_data_row(ws, row, "Export Date/Time:", datetime.now().isoformat(), border)
            
            # Save workbook
            if filename is None:
                filename = f"greeting_{call_session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            file_path = self.output_dir / filename
            wb.save(file_path)
            
            logger.info("Greeting data exported to: %s", file_path)
            return file_path
            
        except Exception as e:
            logger.error("Error exporting greeting data: %s", e)
            raise

    def export_all_calls(
        self,
        filename: str | None = None
    ) -> Path | None:
        """
        Export all calls with greeting data to a single Excel file
        """
        if not OPENPYXL_AVAILABLE:
            logger.error("Excel export not available: openpyxl not installed")
            return None
        
        try:
            calls = self.db.query(CallSession).all()
            
            if not calls:
                logger.warning("No calls to export")
                return None
            
            # Create workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "All Calls"
            
            # Style definitions
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )
            
            # Column widths
            columns = [
                ("Call ID", 12),
                ("Caller Number", 18),
                ("Caller Name", 20),
                ("Device ID", 12),
                ("Language", 15),
                ("Status", 18),
                ("Greeting Played", 18),
                ("Call Date", 20),
                ("Questions Total", 15),
                ("Completion %", 15),
            ]
            
            col_idx = 1
            for col_name, width in columns:
                cell = ws.cell(row=1, column=col_idx)
                cell.value = col_name
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                ws.column_dimensions[cell.column_letter].width = width
                col_idx += 1
            
            # Add data rows
            row = 2
            for call in calls:
                conversation = self.db.query(Conversation).filter(
                    Conversation.call_session_id == call.id
                ).first()
                
                greeting_played = False
                if conversation:
                    greeting_played = any(
                        msg.message_type == "greeting"
                        for msg in conversation.messages
                    )
                
                data = [
                    call.id,
                    call.caller_number,
                    call.caller_name,
                    call.device_id,
                    conversation.language if conversation else "N/A",
                    conversation.status if conversation else "N/A",
                    "Yes" if greeting_played else "No",
                    call.created_at.isoformat() if call.created_at else "N/A",
                    conversation.total_questions if conversation else 0,
                    f"{conversation.completion_percentage}%" if conversation else "N/A",
                ]
                
                for col_idx, value in enumerate(data, 1):
                    cell = ws.cell(row=row, column=col_idx)
                    cell.value = value
                    cell.border = border
                    if isinstance(value, str) and "%" in value:
                        cell.alignment = Alignment(horizontal="center")
                
                row += 1
            
            # Save workbook
            if filename is None:
                filename = f"all_calls_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            file_path = self.output_dir / filename
            wb.save(file_path)
            
            logger.info("All calls data exported to: %s", file_path)
            return file_path
            
        except Exception as e:
            logger.error("Error exporting all calls: %s", e)
            raise

    def _add_section_header(self, ws, row: int, header: str, fill, font, border) -> None:
        """Add a section header row"""
        ws[f"A{row}"] = header
        ws[f"A{row}"].font = font
        ws[f"A{row}"].fill = fill
        ws[f"A{row}"].border = border
        ws.merge_cells(f"A{row}:B{row}")

    def _add_data_row(self, ws, row: int, label: str, value: Any, border) -> None:
        """Add a label-value data row"""
        ws[f"A{row}"] = label
        ws[f"B{row}"] = value
        ws[f"A{row}"].border = border
        ws[f"B{row}"].border = border
        ws[f"A{row}"].font = Font(bold=True)
